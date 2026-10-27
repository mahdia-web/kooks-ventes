"""Analyse la structure des fichiers Excel BROCARD.xlsx et cap frais vente.xlsx"""
from pathlib import Path
from openpyxl import load_workbook
from datetime import datetime, date

FILES = [
    Path('/home/z/my-project/upload/BROCARD.xlsx'),
    Path('/home/z/my-project/upload/cap frais vente.xlsx'),
]

for FILE in FILES:
    print(f"\n{'='*80}")
    print(f"FICHIER : {FILE.name} ({FILE.stat().st_size} octets)")
    print('='*80)

    wb = load_workbook(FILE, data_only=True)
    print(f"Feuilles : {wb.sheetnames}")

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        print(f"\n  --- Feuille : {sheet_name} (max_row={ws.max_row}, max_col={ws.max_column}) ---")

        # Afficher les 8 premières lignes
        print(f"  8 premières lignes :")
        for row_idx in range(1, min(9, ws.max_row + 1)):
            row_data = []
            for col_idx in range(1, min(ws.max_column + 1, 15)):
                v = ws.cell(row=row_idx, column=col_idx).value
                if v is None:
                    row_data.append('None')
                elif isinstance(v, datetime):
                    row_data.append(f"DATE:{v.strftime('%Y-%m-%d')}")
                elif isinstance(v, date):
                    row_data.append(f"date:{v.strftime('%Y-%m-%d')}")
                elif isinstance(v, (int, float)):
                    row_data.append(f"{v}")
                else:
                    s = str(v).strip()
                    if len(s) > 30:
                        s = s[:30] + '...'
                    row_data.append(s)
            print(f"    L{row_idx}: {row_data}")

        # Détecter les en-têtes (recherche des mots clés dans les 5 premières lignes)
        print(f"\n  Détection des en-têtes :")
        for try_row in range(1, 6):
            row_vals = [ws.cell(row=try_row, column=c).value for c in range(1, ws.max_column + 1)]
            row_strs = [str(v).lower().strip() if v else '' for v in row_vals]
            joined = ' | '.join(row_strs)
            if any('date' in s for s in row_strs) and any(
                'montant' in s or 'ca' in s or 'ht' in s or 'enseigne' in s or 'client' in s
                for s in row_strs
            ):
                print(f"    Ligne {try_row} (probable en-têtes) : {row_strs}")
                break
