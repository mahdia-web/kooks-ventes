"""Analyse la structure du fichier Excel SUIVI_AGENTS_FINAL_V2.xlsx.
Liste toutes les feuilles, colonnes, types de données, exemples de valeurs.
"""
from pathlib import Path
from openpyxl import load_workbook

FILE = Path('/home/z/my-project/upload/SUIVI_AGENTS_FINAL_V2.xlsx')

wb = load_workbook(FILE, data_only=True)

print("=" * 80)
print(f"Fichier : {FILE.name}")
print(f"Taille : {FILE.stat().st_size} octets")
print(f"Nombre de feuilles : {len(wb.sheetnames)}")
print(f"Feuilles : {wb.sheetnames}")
print("=" * 80)

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    print(f"\n{'=' * 80}")
    print(f"FEUILLE : {sheet_name}")
    print(f"Dimensions : {ws.dimensions} (max_row={ws.max_row}, max_col={ws.max_column})")
    print(f"{'=' * 80}")

    # Lecture des en-têtes (ligne 1)
    headers = []
    for col_idx in range(1, ws.max_column + 1):
        val = ws.cell(row=1, column=col_idx).value
        headers.append(val)
    print(f"\nEn-têtes ({len(headers)} colonnes) :")
    for i, h in enumerate(headers, 1):
        print(f"  Col {i}: {h!r}")

    # Lecture des 5 premières lignes
    print(f"\n5 premières lignes :")
    for row_idx in range(2, min(7, ws.max_row + 1)):
        row_data = []
        for col_idx in range(1, ws.max_column + 1):
            val = ws.cell(row=row_idx, column=col_idx).value
            row_data.append(val)
        print(f"  Ligne {row_idx}: {row_data}")

    # Détection du type de données par colonne
    print(f"\nType de données par colonne (sur échantillon) :")
    for col_idx in range(1, ws.max_column + 1):
        col_name = ws.cell(row=1, column=col_idx).value
        sample_values = []
        for row_idx in range(2, min(50, ws.max_row + 1)):
            v = ws.cell(row=row_idx, column=col_idx).value
            if v is not None:
                sample_values.append(v)
        if sample_values:
            types = set(type(v).__name__ for v in sample_values[:20])
            print(f"  Col {col_idx} ({col_name}): types={types}, exemples={sample_values[:3]}")
