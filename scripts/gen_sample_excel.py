"""Génère un fichier Excel de ventes d'agents commerciaux pour test."""
import random
from pathlib import Path
from openpyxl import Workbook
from datetime import date, timedelta

output = Path('/home/z/my-project/download/exemple-ventes-agents.xlsx')
output.parent.mkdir(parents=True, exist_ok=True)

wb = Workbook()
ws = wb.active
ws.title = "Ventes 2024"

# En-têtes
ws.append([
    "Agent",
    "Date",
    "Produit",
    "Catégorie",
    "Quantité",
    "Montant",
    "Région",
])

# Données simulées
agents = [
    "Sophie Martin",
    "Thomas Dubois",
    "Léa Bernard",
    "Hugo Petit",
    "Emma Leroy",
    "Luc Moreau",
    "Chloé Garnier",
    "Nathan Roux",
]

products = [
    ("Casque Audio Pro", "Électronique"),
    ("Clavier Mécanique", "Électronique"),
    ("Souris Sans Fil", "Électronique"),
    ("Chaise Ergonomique", "Mobilier"),
    ("Bureau Assis-Debout", "Mobilier"),
    ("Lampe LED", "Mobilier"),
    ("Carnet Premium", "Papeterie"),
    ("Stylo Plume", "Papeterie"),
    ("Tisane Bio", "Alimentation"),
    ("Café Grand Cru", "Alimentation"),
    ("Bouteille inox", "Accessoires"),
    ("Sac à dos urbain", "Accessoires"),
]

regions = ["Paris", "Lyon", "Bordeaux", "Lille", "Marseille", "Nantes"]

start_date = date(2024, 1, 1)
random.seed(42)

for _ in range(150):
    agent = random.choice(agents)
    product, category = random.choice(products)
    qty = random.randint(1, 15)
    unit_price = random.uniform(25, 800)
    amount = round(qty * unit_price, 2)
    days_offset = random.randint(0, 364)
    sale_date = start_date + timedelta(days=days_offset)
    region = random.choice(regions)
    ws.append([agent, sale_date, product, category, qty, amount, region])

# Deuxième feuille avec données Q1 2025
ws2 = wb.create_sheet("Ventes Q1 2025")
ws2.append(["Commercial", "Mois", "Article", "Catégorie", "Volume", "Chiffre d'affaires"])
for _ in range(60):
    agent = random.choice(agents)
    product, category = random.choice(products)
    qty = random.randint(1, 10)
    unit_price = random.uniform(30, 600)
    amount = round(qty * unit_price, 2)
    month = random.choice(["Janvier", "Février", "Mars"])
    ws2.append([agent, month, product, category, qty, amount])

wb.save(output)
print(f"Fichier créé : {output}")
print(f"Taille : {output.stat().st_size} octets")
print(f"Total ventes feuille 1 : {ws.max_row - 1}")
print(f"Total ventes feuille 2 : {ws2.max_row - 1}")
